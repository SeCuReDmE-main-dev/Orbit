"""Render every PDF page and inspect document structure; execute no course code."""
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re

import fitz
from PIL import Image, ImageDraw
from docx import Document
from pypdf import PdfReader

BASE = Path(__file__).resolve().parents[1]
OUTPUT = BASE / 'outputs'
QA = BASE / 'qa'


def review_pdf(path: Path) -> dict:
    source = fitz.open(path)
    reader = PdfReader(path)
    # The exact PDF owns its render directory; old pages cannot impersonate
    # pages of a shorter subsequent edition.
    render = QA / 'render' / (path.stem + '-' + hashlib.sha256(path.read_bytes()).hexdigest()[:12])
    render.mkdir(parents=True, exist_ok=True)
    rows, thumbnails = [], []
    for index, page in enumerate(source):
        pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        image_path = render / f'page-{index + 1:03}.png'
        pixmap.save(image_path)
        spans = [span for block in page.get_text('dict')['blocks'] if 'lines' in block
                 for line in block['lines'] for span in line['spans'] if span.get('text', '').strip()]
        outside = [{'text':span['text'][:140], 'bbox':list(span['bbox'])} for span in spans
                   if span['bbox'][0] < 0 or span['bbox'][1] < 0
                   or span['bbox'][2] > page.rect.width or span['bbox'][3] > page.rect.height]
        near_edge = [{'text':span['text'][:140], 'bbox':list(span['bbox'])} for span in spans
                     if span['bbox'][0] < 25 or span['bbox'][2] > page.rect.width - 25]
        overlaps = []
        for position, left in enumerate(spans):
            x0, y0, x1, y1 = left['bbox']
            for right in spans[position + 1:]:
                rx0, ry0, rx1, ry1 = right['bbox']
                vertical = min(y1, ry1) - max(y0, ry0)
                horizontal = min(x1, rx1) - max(x0, rx0)
                if horizontal > 2 and vertical > min(y1 - y0, ry1 - ry0) * .7:
                    overlaps.append({'left':left['text'][:100], 'right':right['text'][:100],
                                     'overlapPt':round(horizontal, 2), 'review':'candidate-not-an-automatic-defect'})
        text = page.get_text()
        rows.append({'page':index + 1, 'textCharacters':len(text), 'links':len(page.get_links()),
                     'outsidePage':outside, 'nearHorizontalEdge':near_edge, 'overlapCandidates':overlaps,
                     'replacementCharacters':text.count('\ufffd'),
                     'render':image_path.relative_to(BASE).as_posix(),
                     'pngSha256':hashlib.sha256(image_path.read_bytes()).hexdigest()})
        thumbnail = Image.open(image_path).convert('RGB')
        thumbnail.thumbnail((300, 425))
        thumbnails.append(thumbnail)
    sheets = []
    for start in range(0, len(thumbnails), 12):
        sheet = Image.new('RGB', (4 * 330, 3 * 465), '#dce3e8')
        draw = ImageDraw.Draw(sheet)
        for item, thumbnail in enumerate(thumbnails[start:start + 12]):
            x, y = (item % 4) * 330, (item // 4) * 465
            draw.text((x + 12, y + 5), f'Page {start + item + 1}', fill='#18212b')
            sheet.paste(thumbnail, (x + 12, y + 25))
        sheet_path = render / f'contact-{start // 12 + 1:02}.png'
        sheet.save(sheet_path)
        sheets.append(sheet_path.relative_to(BASE).as_posix())
    metadata = dict(reader.metadata or {})
    return {'path':path.relative_to(BASE).as_posix(), 'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
            'pages':len(source), 'allPagesRendered':True, 'metadata':metadata,
            'authorMatches':reader.metadata.author == 'Jean-Sébastien Beaulieu',
            'language':reader.trailer['/Root'].get('/Lang'),
            'textCharacters':sum(row['textCharacters'] for row in rows),
            'totalLinks':sum(row['links'] for row in rows),
            'outsidePageSpanCount':sum(len(row['outsidePage']) for row in rows),
            'overlapCandidateCount':sum(len(row['overlapCandidates']) for row in rows),
            'replacementCharacterCount':sum(row['replacementCharacters'] for row in rows),
            'pageRows':rows, 'contactSheets':sheets,
            'visualInspection':'PENDING', 'semanticLinkReadback':'PENDING'}


def review_docx(path: Path) -> dict:
    source = Document(path)
    text = '\n'.join(paragraph.text for paragraph in source.paragraphs)
    cells = '\n'.join(cell.text for table in source.tables for row in table.rows for cell in row.cells)
    return {'path':path.relative_to(BASE).as_posix(), 'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
            'title':source.core_properties.title, 'author':source.core_properties.author,
            'authorMatches':source.core_properties.author == 'Jean-Sébastien Beaulieu',
            'paragraphs':len(source.paragraphs), 'tables':len(source.tables),
            'headings':sum(paragraph.style.name.startswith('Heading') for paragraph in source.paragraphs),
            'bodyCharacters':len(text), 'tableCharacters':len(cells),
            'formatTotal40Hours':bool(re.search(r'40\s*heures', text)),
            'canonicalDisclosureCount':text.replace('’', "'").count("J'ai utilisé Codex à son plein potentiel comme partenaire de recherche"),
            'layoutInspection':'PENDING_NATIVE_RENDER'}


if __name__ == '__main__':
    QA.mkdir(exist_ok=True)
    report = {'state':'OUTPUT_STRUCTURE_AND_ALL_PAGE_RENDER_OBSERVED',
              'recordedAt':datetime.now(timezone.utc).isoformat(),
              'courseSoftwareExecuted':False, 'learnerUnderstandingExamined':False,
              'readiness':'review-ready-provisional', 'deliveryReady':False, 'fullMissionComplete':False,
              'pdf':[review_pdf(path) for path in sorted(OUTPUT.glob('*.pdf'))],
              'docx':[review_docx(path) for path in sorted(OUTPUT.glob('*.docx')) if path.name != 'reference.docx'],
              'limits':['Rendered pages still require visual review.',
                        'DOCX structure is distinct from native Word pagination.',
                        'External links need separate readback.',
                        'Course Context admission HOLD; Canva not produced.',
                        'Assembled browser and authenticated Studio retain separate evidence.']}
    (QA / 'output-review.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8', newline='\n')
    print(json.dumps({'state':report['state'],
                      'pdf':[{key:row[key] for key in ['path','pages','allPagesRendered','outsidePageSpanCount','overlapCandidateCount','replacementCharacterCount','authorMatches','language']} for row in report['pdf']],
                      'docx':[{key:row[key] for key in ['path','paragraphs','tables','authorMatches','canonicalDisclosureCount']} for row in report['docx']]}, ensure_ascii=False))
